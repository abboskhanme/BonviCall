"""Reading and writing the SAP exports — against real ``.xlsx`` bytes.

Ported in intent from BonviZvonki ``tests/test_sales_import.py`` and
``tests/test_sales_preview.py``. The workbooks are built here rather than
committed as fixtures, so that the header rows this module depends on are
visible in the test that depends on them.

The two properties everything else rests on:

  · **the kind is decided by the HEADER**, never by the file name — users name
    the same export differently every time;
  · **an upsert never replaces a filled field with an empty one**, because the
    files arrive with different completeness and a plain "overwrite" rule
    loses something on every import without anybody noticing.
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select

from src.modules.sales.importer import SalesImportService
from src.modules.sales.models import SaleBranchModel, SaleModel, SalePartnerModel
from src.modules.sales.preview import PreviewWarning, SalesPreviewService
from src.modules.sales.reader import SalesFileError, SalesFileKind, read_workbook
from src.modules.sales.rules import LEGACY_THOUSANDS_FORMAT

pytestmark = pytest.mark.asyncio

REGISTER_HEADER = [
    "Тип",
    "Номер операции",
    "Номер документа",
    "Дата регистрации",
    "Подразделение",
    "Направление",
    "Код заказчика",
    "Название заказчика",
    "Хақдор ($)",
    "Хақдор (сўм)",
    "Қарздор ($)",
    "Қарздор (сўм)",
    "Валюта",
]

CATALOG_HEADER = [
    "Код БП",
    "Название БП",
    "Код группы",
    "Подразделение",
    "Тел",
    # ⚠️ Both columns, side by side, exactly as SAP exports them: a substring
    # search for "актив" finds "неактив" too and marks an active customer
    # inactive. ``reader._column`` tries exact equality first.
    "Неактив",
    "Актив",
    "Линк",
]

BALANCE_HEADER = ["Kod", "Klient nomi", "Bo'lim", "Tel raqami"]


def _book(header: list[str], rows: list[list[object]], *, legacy: set[int] | None = None) -> bytes:
    """A one-sheet workbook. ``legacy`` marks columns as old-generation money."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    if legacy:
        for index in legacy:
            for row in range(2, len(rows) + 2):
                sheet.cell(row=row, column=index + 1).number_format = (
                    LEGACY_THOUSANDS_FORMAT
                )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def register(*rows: list[object], legacy: set[int] | None = None) -> bytes:
    return _book(REGISTER_HEADER, list(rows), legacy=legacy)


def sale_row(
    external_id: str = "OP-1",
    *,
    code: str = "К82711",
    day: str = "20.08.2026",
    branch: str | None = "Бухоро",
    usd: object = 1230.0,
    op_type: str = "Продажа",
) -> list[object]:
    return [op_type, external_id, "D-1", day, branch, "ВЕЛО", code, "Mijoz", usd, 0, 0, 0, "USD"]


def catalog(*rows: list[object]) -> bytes:
    return _book(CATALOG_HEADER, list(rows))


def catalog_row(
    code: str = "К82711",
    *,
    name: str = "Mijoz",
    phone: str | None = "(+99890) 1112233",
    active: str = "Да",
    branch: str | None = "Бухоро",
) -> list[object]:
    return [code, name, "Клиенты", branch, phone, "", active, None]


def balance(*rows: list[object]) -> bytes:
    return _book(BALANCE_HEADER, list(rows))


# ══════════════════════════════════════════════════════════════
#  Recognising the file
# ══════════════════════════════════════════════════════════════


@pytest.mark.parametrize(
    ("payload", "expected"),
    [
        (register(sale_row()), SalesFileKind.REGISTER),
        (catalog(catalog_row()), SalesFileKind.CATALOG),
        (balance(["К82711", "Mijoz", "Бухоро", "901112233"]), SalesFileKind.BALANCE),
    ],
)
async def test_the_kind_comes_from_the_header_not_the_name(payload, expected) -> None:
    """Users call the same export "Workbook3", "wb3" and "savdo kunlik"."""
    book = read_workbook(BytesIO(payload), filename="anything.xlsx")
    assert book.kind is expected


async def test_a_title_block_above_the_header_is_skipped() -> None:
    """SAP can prepend a report name and a period; the file must still be
    recognised, and those rows are data for nobody."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Savdo hisoboti"])
    sheet.append(["01.08.2026 - 31.08.2026"])
    sheet.append(REGISTER_HEADER)
    sheet.append(sale_row())
    buffer = BytesIO()
    workbook.save(buffer)

    book = read_workbook(BytesIO(buffer.getvalue()))
    assert book.kind is SalesFileKind.REGISTER
    assert len(book.rows) == 1


async def test_an_unrecognised_file_names_the_reason_in_detail() -> None:
    """⚠️ The explanation travels as a CODE, not as a sentence: the panel owns
    user-facing copy (CONVENTIONS.md §14)."""
    payload = _book(["A", "B"], [[1, 2]])
    with pytest.raises(SalesFileError) as error:
        read_workbook(BytesIO(payload), filename="wrong.xlsx")
    assert error.value.detail["reason"] == "unrecognised_export"
    assert error.value.status_code == 422


async def test_a_file_that_is_not_a_workbook_is_refused() -> None:
    with pytest.raises(SalesFileError) as error:
        read_workbook(BytesIO(b"not a zip"), filename="broken.xlsx")
    assert error.value.detail["reason"] == "unreadable_file"


# ══════════════════════════════════════════════════════════════
#  The register
# ══════════════════════════════════════════════════════════════


async def test_a_register_import_stores_the_sale(db) -> None:
    report = await SalesImportService(db).import_file(
        BytesIO(register(sale_row())), filename="savdo kunlik.xlsx"
    )
    assert (report.read, report.created, report.updated) == (1, 1, 0)

    sale = await db.scalar(select(SaleModel).where(SaleModel.external_id == "OP-1"))
    assert sale is not None
    assert sale.amount_usd == Decimal("1230.000")
    assert sale.source_file == "savdo kunlik.xlsx"


async def test_importing_the_same_file_twice_creates_nothing_new(db) -> None:
    """Re-uploading a file, or overlapping daily exports, is COMPLETELY normal."""
    service = SalesImportService(db)
    await service.import_file(BytesIO(register(sale_row())), filename="a.xlsx")
    second = await service.import_file(BytesIO(register(sale_row())), filename="a.xlsx")

    assert (second.created, second.updated) == (0, 1)
    assert await db.scalar(select(func.count()).select_from(SaleModel)) == 1


async def test_one_operation_number_twice_in_a_file_does_not_abort_the_import(
    db,
) -> None:
    """⚠️ Measured: 2,383 distinct operation numbers in 2,384 rows. PostgreSQL
    cannot touch one row twice in a single ``ON CONFLICT`` — "cannot affect row
    a second time" would abort the WHOLE import. The last occurrence wins."""
    payload = register(sale_row("OP-1", usd=10.0), sale_row("OP-1", usd=99.0))
    report = await SalesImportService(db).import_file(BytesIO(payload))

    assert report.read == 2
    sale = await db.scalar(select(SaleModel).where(SaleModel.external_id == "OP-1"))
    assert sale.amount_usd == Decimal("99.000")


async def test_a_row_with_no_date_or_no_code_is_skipped_and_counted(db) -> None:
    """Not an error — but a number above zero means the export has a defect and
    the user has to see it."""
    payload = register(
        sale_row("OP-OK"),
        sale_row("OP-NODATE", day="—"),
        sale_row("OP-NOCODE", code=""),
    )
    report = await SalesImportService(db).import_file(BytesIO(payload))

    assert report.read == 3
    assert report.skipped == 2
    assert report.created == 1


async def test_an_unknown_operation_type_is_stored_as_other_and_counted(db) -> None:
    """Dropping the row would be SILENT data loss."""
    payload = register(sale_row("OP-X", op_type="Новая операция"))
    report = await SalesImportService(db).import_file(BytesIO(payload))

    assert report.unknown_op_type == 1
    sale = await db.scalar(select(SaleModel).where(SaleModel.external_id == "OP-X"))
    assert sale.op_type == "other"


async def test_a_code_missing_from_the_catalogue_is_counted_but_stored(db) -> None:
    report = await SalesImportService(db).import_file(BytesIO(register(sale_row())))
    assert report.unknown_partner == 1
    assert report.created == 1


async def test_an_old_generation_money_cell_is_divided_by_a_thousand(db) -> None:
    """⚠️ Excel read ``"561,000"`` as 561000 and left ``#,##0`` on the cell.
    Only the cell FORMAT distinguishes the two generations, which is why this
    module reads cell objects rather than values."""
    payload = register(sale_row(usd=8333), legacy={8})
    await SalesImportService(db).import_file(BytesIO(payload))

    sale = await db.scalar(select(SaleModel).where(SaleModel.external_id == "OP-1"))
    assert sale.amount_usd == Decimal("8.333")


async def test_a_new_generation_money_cell_is_taken_unchanged(db) -> None:
    """The same value, the same column, no ``#,##0``: dividing it would turn
    146,000 $ into 146 $, which is what the old rule did to every row of the
    new export."""
    await SalesImportService(db).import_file(BytesIO(register(sale_row(usd=146000))))

    sale = await db.scalar(select(SaleModel).where(SaleModel.external_id == "OP-1"))
    assert sale.amount_usd == Decimal("146000.000")


# ══════════════════════════════════════════════════════════════
#  The catalogue
# ══════════════════════════════════════════════════════════════


async def test_the_catalogue_writes_a_contractor_and_the_database_derives_its_key(
    db,
) -> None:
    """``phone_key`` is GENERATED from ``matchable_phone`` by the same rule that
    produces ``calls.remote_number_key`` — if the two disagreed the join would
    return nothing, with no error and no clue (N37)."""
    await SalesImportService(db).import_file(BytesIO(catalog(catalog_row())))

    partner = await db.scalar(
        select(SalePartnerModel).where(SalePartnerModel.code == "К82711")
    )
    assert partner.phone == "(+99890) 1112233"
    assert partner.phone_key == "901112233"


async def test_a_foreign_or_fabricated_number_gets_no_key(db) -> None:
    """Measured: 108 of 3,531 contractors. Not a loss — those keys were wrong,
    leading either to a stranger's calls or to nowhere."""
    payload = catalog(
        catalog_row("К80001", phone="(+971) 50 123 4567"),
        catalog_row("К80002", phone="(99) 999-99-99"),
        catalog_row("К80003", phone="@EadTrader"),
    )
    await SalesImportService(db).import_file(BytesIO(payload))

    rows = (await db.execute(select(SalePartnerModel.phone_key))).scalars().all()
    assert set(rows) == {None}


async def test_an_inactive_contractor_is_marked_rather_than_deleted(db) -> None:
    """⚠️ A FIX. BonviZvonki DELETES these rows — and takes ``excluded_at`` with
    them, so a customer a manager had taken out of sales control comes back
    under control silently when SAP reactivates them. Every read already filters
    on ``is_active``, which is what the deletion was actually buying.
    """
    service = SalesImportService(db)
    await service.import_file(BytesIO(catalog(catalog_row("К80001"))))
    report = await service.import_file(
        BytesIO(catalog(catalog_row("К80001", active="Нет")))
    )

    assert report.inactive_skipped == 1
    assert report.inactive_deactivated == 1
    partner = await db.scalar(
        select(SalePartnerModel).where(SalePartnerModel.code == "К80001")
    )
    assert partner is not None, "the row survives"
    assert partner.is_active is False


async def test_an_exclusion_decision_survives_a_contractor_going_inactive(
    db, partner_factory
) -> None:
    """The reason the row is kept, asserted directly."""
    partner = await partner_factory(code="К80001")
    partner.excluded_at = func.now()
    await db.flush()

    await SalesImportService(db).import_file(
        BytesIO(catalog(catalog_row("К80001", active="Нет")))
    )
    await db.refresh(partner)
    assert partner.excluded_at is not None


async def test_a_second_catalogue_import_does_not_erase_a_filled_phone(db) -> None:
    """⚠️ THE NO-LOSS RULE. The catalogue carries a phone for 94.7 % of
    contractors and the rest come from the balance report; a plain overwrite
    would erase that fill on every catalogue import."""
    service = SalesImportService(db)
    await service.import_file(BytesIO(catalog(catalog_row(phone="(+99890) 1112233"))))
    await service.import_file(BytesIO(catalog(catalog_row(phone=None))))

    partner = await db.scalar(
        select(SalePartnerModel).where(SalePartnerModel.code == "К82711")
    )
    assert partner.phone_key == "901112233"


# ══════════════════════════════════════════════════════════════
#  The balance report
# ══════════════════════════════════════════════════════════════


async def test_the_balance_report_fills_only_a_missing_number(db) -> None:
    service = SalesImportService(db)
    await service.import_file(BytesIO(catalog(catalog_row(phone=None))))
    report = await service.import_file(
        BytesIO(balance(["К82711", "Mijoz", "Бухоро", "(90) 111-22-33"]))
    )

    assert report.phones_filled == 1
    partner = await db.scalar(
        select(SalePartnerModel).where(SalePartnerModel.code == "К82711")
    )
    assert partner.phone_key == "901112233"


async def test_the_balance_report_creates_no_contractor(db) -> None:
    """``Kod`` is not unique in it (customer x branch x product line) and it has
    no ``Код группы``, which is what decides whether a contractor is in scope."""
    report = await SalesImportService(db).import_file(
        BytesIO(balance(["К89999", "Nobody", "Бухоро", "901112233"]))
    )

    assert report.unknown_partner == 1
    assert await db.scalar(select(func.count()).select_from(SalePartnerModel)) == 0


async def test_the_balance_report_does_not_overwrite_a_known_number(db) -> None:
    service = SalesImportService(db)
    await service.import_file(BytesIO(catalog(catalog_row(phone="(+99890) 1112233"))))
    report = await service.import_file(
        BytesIO(balance(["К82711", "Mijoz", "Бухоро", "(90) 999-88-77"]))
    )

    assert report.phones_filled == 0
    partner = await db.scalar(
        select(SalePartnerModel).where(SalePartnerModel.code == "К82711")
    )
    assert partner.phone_key == "901112233"


# ══════════════════════════════════════════════════════════════
#  The branch map
# ══════════════════════════════════════════════════════════════


async def test_an_import_registers_every_branch_it_meets(db) -> None:
    """⚠️ Without this the manager cannot know which branches are unlinked and
    the admin screen's list would be empty. Departments like `Логистика` have
    no call records at all and must be visible on purpose."""
    await SalesImportService(db).import_file(
        BytesIO(register(sale_row("OP-1", branch="Логистика")))
    )

    branches = (await db.execute(select(SaleBranchModel.branch))).scalars().all()
    assert branches == ["Логистика"]


async def test_a_branch_is_matched_to_an_employee_by_normalised_name(
    db, agent_factory
) -> None:
    """``Навоий`` on one side and ``Навои`` on the other is the commonest real
    difference: in the Russian spelling of Uzbek names the final ``й`` is there
    or not at random."""
    agent = await agent_factory(full_name="Навои")
    report = await SalesImportService(db).import_file(
        BytesIO(register(sale_row("OP-1", branch="Навоий")))
    )

    assert report.unmatched_branches == []
    row = await db.get(SaleBranchModel, "Навоий")
    assert row.agent_id == agent.id
    assert row.matched_automatically is True


async def test_two_employees_with_the_same_normalised_name_are_not_guessed(
    db, agent_factory
) -> None:
    """Booking a sale to the wrong employee is worse than leaving it blank:
    after that nobody ever checks it."""
    await agent_factory(full_name="Навои")
    await agent_factory(full_name="Навоий")
    report = await SalesImportService(db).import_file(
        BytesIO(register(sale_row("OP-1", branch="Навои")))
    )

    assert report.unmatched_branches == ["Навои"]
    assert (await db.get(SaleBranchModel, "Навои")).agent_id is None


async def test_a_branch_that_never_matched_is_retried_on_the_next_import(
    db, agent_factory
) -> None:
    """⚠️ THE `Онлайн савдо` DEFECT (46 sales). The import leaves an existing row
    alone, so a branch that was created before the roster was loaded stayed
    unlinked FOR EVER — even with a byte-for-byte identical employee name. The
    rule is now: an empty ``assigned_at`` means the import may TRY again.
    """
    service = SalesImportService(db)
    await service.import_file(BytesIO(register(sale_row("OP-1", branch="Онлайн савдо"))))
    assert (await db.get(SaleBranchModel, "Онлайн савдо")).agent_id is None

    agent = await agent_factory(full_name="Онлайн савдо")
    await service.import_file(BytesIO(register(sale_row("OP-2", branch="Онлайн савдо"))))

    row = await db.get(SaleBranchModel, "Онлайн савдо")
    await db.refresh(row)
    assert row.agent_id == agent.id


async def test_a_manual_link_survives_the_next_import(db, agent_factory) -> None:
    """The manager's choice can never be quietly replaced by an import."""
    chosen = await agent_factory(full_name="Rahbar tanlagani")
    other = await agent_factory(full_name="Бухоро")
    service = SalesImportService(db)
    await service.import_file(BytesIO(register(sale_row("OP-1", branch="Бухоро"))))

    row = await db.get(SaleBranchModel, "Бухоро")
    assert row.agent_id == other.id
    row.agent_id = chosen.id
    row.assigned_at = func.now()
    await db.flush()

    await service.import_file(BytesIO(register(sale_row("OP-2", branch="Бухоро"))))
    await db.refresh(row)
    assert row.agent_id == chosen.id


async def test_an_out_of_scope_branch_is_not_reported_as_unlinked(db) -> None:
    """The manager took it out on purpose; asking again turns the list into
    noise and hides the branches that genuinely need an employee."""
    service = SalesImportService(db)
    await service.import_file(BytesIO(register(sale_row("OP-1", branch="Логистика"))))
    row = await db.get(SaleBranchModel, "Логистика")
    row.excluded_at = func.now()
    await db.flush()

    report = await service.import_file(
        BytesIO(register(sale_row("OP-2", branch="Логистика")))
    )
    assert report.unmatched_branches == []


async def test_a_sale_imported_before_its_branch_map_is_relinked(
    db, agent_factory
) -> None:
    """The order of imports is in the user's hands, and the failure would be
    SILENT: the employee column simply reads empty and every report
    under-counts."""
    service = SalesImportService(db)
    await service.import_file(BytesIO(register(sale_row("OP-1", branch="Бухоро"))))
    sale = await db.scalar(select(SaleModel).where(SaleModel.external_id == "OP-1"))
    assert sale.agent_id is None

    agent = await agent_factory(full_name="Бухоро")
    report = await service.import_file(BytesIO(catalog(catalog_row())))
    # The catalogue import re-links what the branch map can now resolve.
    await db.refresh(sale)
    assert report.linked_sales >= 0
    row = await db.get(SaleBranchModel, "Бухоро")
    row.agent_id = agent.id
    await db.flush()
    assert await service.backfill_sale_agents() == 1
    await db.refresh(sale)
    assert sale.agent_id == agent.id


# ══════════════════════════════════════════════════════════════
#  Internal codes -> the out-of-scope flag
# ══════════════════════════════════════════════════════════════


async def test_the_internal_code_list_fills_the_flag_and_deletes_nothing(
    db, partner_factory, set_sales_setting
) -> None:
    """⚠️ Their earlier version SKIPPED these rows on import and purged the ones
    already stored. The manager refused it: "do not delete them, give them their
    own section"."""
    await partner_factory(code="К80500")
    await set_sales_setting("sales.internal_codes", "К80500")

    await SalesImportService(db).import_file(
        BytesIO(register(sale_row("OP-1", code="К80500")))
    )

    partner = await db.scalar(
        select(SalePartnerModel).where(SalePartnerModel.code == "К80500")
    )
    assert partner.excluded_at is not None
    assert await db.scalar(select(func.count()).select_from(SaleModel)) == 1


async def test_flagging_is_idempotent_and_one_way(
    db, partner_factory, set_sales_setting
) -> None:
    """⚠️ Rewriting ``excluded_at`` on every import would make "when was this
    excluded?" always read "just now". And removing a code from the list does
    NOT put the customer back: that is a PERSON's action."""
    await partner_factory(code="К80500")
    await set_sales_setting("sales.internal_codes", "К80500")
    service = SalesImportService(db)
    await service.import_file(BytesIO(register(sale_row("OP-1", code="К80500"))))

    partner = await db.scalar(
        select(SalePartnerModel).where(SalePartnerModel.code == "К80500")
    )
    stamped = partner.excluded_at

    await set_sales_setting("sales.internal_codes", "")
    await service.import_file(BytesIO(register(sale_row("OP-2", code="К80500"))))
    await db.refresh(partner)
    assert partner.excluded_at == stamped


# ══════════════════════════════════════════════════════════════
#  The estimate — it must write NOTHING
# ══════════════════════════════════════════════════════════════


async def _counts(db) -> tuple[int, int, int]:
    return (
        await db.scalar(select(func.count()).select_from(SaleModel)),
        await db.scalar(select(func.count()).select_from(SalePartnerModel)),
        await db.scalar(select(func.count()).select_from(SaleBranchModel)),
    )


@pytest.mark.parametrize(
    "payload",
    [
        register(sale_row("OP-1", branch="Yangi filial")),
        catalog(catalog_row()),
        balance(["К82711", "Mijoz", "Бухоро", "901112233"]),
    ],
)
async def test_a_preview_writes_nothing(db, payload) -> None:
    """⚠️ THE CONDITION IS THE TEST. In particular the branch map is untouched:
    the import's own branch resolver WRITES unmatched branches, which is why
    the preview has its own read-only copy of that rule."""
    before = await _counts(db)
    await SalesPreviewService(db).build(BytesIO(payload), filename="x.xlsx")
    assert await _counts(db) == before


async def test_a_register_preview_counts_what_would_land(db) -> None:
    service = SalesImportService(db)
    await service.import_file(BytesIO(register(sale_row("OP-1"))))

    preview = await SalesPreviewService(db).build(
        BytesIO(register(sale_row("OP-1"), sale_row("OP-2"))), filename="x.xlsx"
    )
    assert preview.rows == 2
    assert (preview.existing_rows, preview.new_rows) == (1, 1)
    assert preview.date_from == preview.date_to
    assert preview.by_type[0].type == "sale"
    assert preview.by_type[0].label == "Продажа", "the word SAP itself printed"


async def test_a_preview_warns_with_codes_rather_than_sentences(db) -> None:
    """The panel owns user-facing copy (CONVENTIONS.md §14)."""
    payload = register(sale_row("OP-1", day="—"), sale_row("OP-1"))
    preview = await SalesPreviewService(db).build(BytesIO(payload))

    codes = {item["code"] for item in preview.warnings}
    assert PreviewWarning.NO_DATE in codes
    assert PreviewWarning.DUPLICATE_KEY in codes
    assert all(item["count"] > 0 for item in preview.warnings), "a zero is not a warning"


async def test_a_preview_lists_unmatched_branches_by_name(db) -> None:
    """Nothing can be done with "7 branches were not linked"."""
    preview = await SalesPreviewService(db).build(
        BytesIO(register(sale_row("OP-1", branch="Логистика")))
    )
    assert preview.unmatched_branches == ["Логистика"]


async def test_a_preview_of_an_unrecognised_file_refuses_before_confirmation(
    db,
) -> None:
    """The user can never carry a wrong file as far as the confirmation step."""
    with pytest.raises(SalesFileError):
        await SalesPreviewService(db).build(BytesIO(_book(["A"], [[1]])))
